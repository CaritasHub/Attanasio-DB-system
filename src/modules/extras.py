from flask import Blueprint, request, send_file, jsonify
from db import get_db_connection
from .utils import login_required, role_required
from openpyxl import Workbook
import csv
import io

# Mapping from endpoint names to database table names
TABLE_MAP = {
    'specialists': 'Specialista',
    'users': 'Utente',
    'afferenze': 'Afferenza',
    'sedi': 'Sede',
    'provvedimenti': 'Provvedimento',
}

extras_bp = Blueprint('extras', __name__, url_prefix='/extras')


@extras_bp.route('/export/<name>', methods=['GET'])
@login_required
def export_table(name):
    table = TABLE_MAP.get(name)
    if not table:
        return jsonify({'error': 'unknown table'}), 400
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    cur.execute(f'SELECT * FROM {table}')
    rows = cur.fetchall()
    headers = []
    if rows:
        headers = list(rows[0].keys())
    else:
        cur.close()
        cur = conn.cursor()
        cur.execute(f"SHOW COLUMNS FROM {table}")
        headers = [row[0] for row in cur.fetchall()]
    cur.close()

    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'{table}.xlsx'
    return send_file(output, download_name=filename,
                     as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@extras_bp.route('/import/<name>', methods=['POST'])
@login_required
def import_csv(name):
    table = TABLE_MAP.get(name)
    if not table:
        return jsonify({'error': 'unknown table'}), 400
    file = request.files.get('file')
    if not file:
        return jsonify({'error': 'missing file'}), 400

    stream = io.StringIO(file.stream.read().decode('utf-8'))
    reader = csv.DictReader(stream)
    columns = reader.fieldnames
    if not columns:
        return jsonify({'error': 'empty csv'}), 400

    conn = get_db_connection()
    cur = conn.cursor()
    placeholders = ','.join(['%s'] * len(columns))
    columns_sql = ','.join(columns)
    query = f'INSERT INTO {table} ({columns_sql}) VALUES ({placeholders})'
    try:
        conn.start_transaction()
        for row in reader:
            values = [row.get(c) or None for c in columns]
            cur.execute(query, values)
        conn.commit()
    except Exception as e:
        conn.rollback()
        cur.close()
        return jsonify({'error': str(e)}), 400
    cur.close()
    return jsonify({'status': 'imported'})


@extras_bp.route('/columns/<name>', methods=['GET'])
@login_required
def get_columns(name):
    """Return the column names for a given table."""
    table = TABLE_MAP.get(name)
    if not table:
        return jsonify({'error': 'unknown table'}), 400

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute(f"SHOW COLUMNS FROM {table}")
    columns = [row[0] for row in cur.fetchall()]
    cur.close()
    return jsonify(columns)


@extras_bp.route('/column-config/<name>', methods=['GET'])
@login_required
def get_column_config(name):
    table = TABLE_MAP.get(name)
    if not table:
        return jsonify({'error': 'unknown table'}), 400
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)

    cur.execute(
        'SELECT column_name, visible, display_order FROM ColumnConfig '
        'WHERE table_name=%s ORDER BY display_order',
        (table,)
    )

    rows = cur.fetchall()
    cur.close()
    return jsonify({'columns': rows})


@extras_bp.route('/column-config/<name>', methods=['POST'])
@login_required
@role_required('editor', 'founder')
def set_column_config(name):
    table = TABLE_MAP.get(name)
    if not table:
        return jsonify({'error': 'unknown table'}), 400
    data = request.json or {}
    columns = data.get('columns', [])
    conn = get_db_connection()
    cur = conn.cursor()
    conn.start_transaction()
    try:
        cur.execute('DELETE FROM ColumnConfig WHERE table_name=%s', (table,))
        for idx, col in enumerate(columns):
            cur.execute(
                'INSERT INTO ColumnConfig (table_name, column_name, visible, display_order) VALUES (%s,%s,%s,%s)',
                (table, col.get('column_name'), bool(col.get('visible', True)), int(col.get('display_order', idx)))
            )
        conn.commit()
    except Exception as e:
        conn.rollback()
        cur.close()
        return jsonify({'error': str(e)}), 400
    cur.close()
    return jsonify({'status': 'saved'})
